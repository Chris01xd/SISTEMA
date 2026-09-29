"""/api/v1/ml/*: Backend <-> Langflow contract and the direct (in-process) engine.

The Langflow server is replaced by an httpx MockTransport whose flow output is a REAL
analysis result produced by ai_module.ml on synthetic MIMIC-shaped data."""
import json
import uuid
from types import SimpleNamespace

import httpx
import pytest
from fastapi.testclient import TestClient

pytest.importorskip("sklearn")
pytest.importorskip("langchain_core")

from app.core.config import settings  # noqa: E402
from app.services import ai_module_service, ml_analysis_service  # noqa: E402
from app.services.langflow_client import LangflowClient  # noqa: E402

ai_module_service._import("ai_module")  # repo root on sys.path

import ai_module.ml.service as ml_service_module  # noqa: E402
from ai_module.config.settings import AISettings  # noqa: E402
from ai_module.digital_twin.service import DigitalTwinService  # noqa: E402
from ai_module.ml.config import MLSettings  # noqa: E402
from ai_module.tests.conftest import make_raw_tables  # noqa: E402

MODELS = ["logistic_regression", "svm", "random_forest", "gradient_boosting", "voting_ensemble", "stacking_ensemble"]


@pytest.fixture(scope="module")
def ml_env(tmp_path_factory):
    root = tmp_path_factory.mktemp("mlapi")
    ai = AISettings(_env_file=None, data_dir=root / "data", artifacts_dir=root / "artifacts", llm_provider="none")
    ml = MLSettings(_env_file=None, runs_dir=root / "artifacts" / "ml_runs", explain_with_llm=False, bootstrap_resamples=100)
    DigitalTwinService(ai).preprocess_data(make_raw_tables(n_patients=60))
    output = ml_service_module.MLAnalysisService(ai, ml).run({"models": MODELS})
    return SimpleNamespace(ai=ai, ml=ml, output=output)


@pytest.fixture
def env(monkeypatch, ml_env):
    from app.authentication.dependencies import get_current_user
    from app.main import app

    admin = SimpleNamespace(id=uuid.uuid4(), role=SimpleNamespace(permissions={"*": True}), is_active=True)
    monkeypatch.setitem(app.dependency_overrides, get_current_user, lambda: admin)
    monkeypatch.setattr(ml_service_module, "get_settings", lambda: ml_env.ai)
    monkeypatch.setattr(ml_service_module, "get_ml_settings", lambda: ml_env.ml)
    monkeypatch.setattr(settings, "ML_ANALYSIS_MODE", "langflow")
    monkeypatch.setattr(settings, "LANGFLOW_API_KEY", "lf-test-key")
    monkeypatch.setattr(settings, "LANGFLOW_FLOW_ID", "ml-model-comparison")
    return SimpleNamespace(client=TestClient(app), app=app, output=ml_env.output)


def _langflow(monkeypatch, handler):
    seen: list[httpx.Request] = []

    def record(request: httpx.Request) -> httpx.Response:
        seen.append(request)
        return handler(request)

    monkeypatch.setattr(ml_analysis_service, "client",
                        lambda: LangflowClient("http://langflow:7860", settings.LANGFLOW_API_KEY, 30, transport=httpx.MockTransport(record)))
    return seen


def _run_response(text: str) -> dict:
    return {"session_id": "s", "outputs": [{"inputs": {"input_value": "x"}, "outputs": [
        {"results": {"message": {"text": text, "sender": "Machine"}}, "artifacts": {}, "outputs": {}}]}]}


def test_analysis_runs_through_langflow_and_returns_the_validated_result(env, monkeypatch):
    seen = _langflow(monkeypatch, lambda r: httpx.Response(200, json=_run_response(json.dumps(env.output))))
    r = env.client.post("/api/v1/ml/analysis", json={"dataset": "mimic_icu", "target": "icu_flag", "cv_folds": 5})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["engine"] == "langflow" and body["langflow"]["flow"] == "ml-model-comparison"
    assert body["best_model"]["name"] == env.output["ranking"][0]["model"]
    assert body["models_evaluated"] == 6 and len(body["hybrid_models"]) == 2 and len(body["normal_models"]) == 4
    # Nothing the frontend renders is dropped by the response model.
    for key in ("eda", "statistical_tests", "hyperparameters", "curves", "interpretations", "crisp_dm"):
        assert body[key] == env.output[key], key
    # What Langflow received: run API contract, parameters as JSON in input_value, API key header.
    request = seen[0]
    assert request.url.path == "/api/v1/run/ml-model-comparison" and request.url.params["stream"] == "false"
    assert request.headers["x-api-key"] == "lf-test-key"
    sent = json.loads(request.content)
    assert sent["input_type"] == "chat" and sent["output_type"] == "chat"
    assert json.loads(sent["input_value"]) == {"dataset": "mimic_icu", "target": "icu_flag", "problem_type": "classification",
                                               "cv_folds": 5, "language": "es"}


def test_langflow_unreachable_is_503(env, monkeypatch):
    def refuse(request):
        raise httpx.ConnectError("connection refused")

    _langflow(monkeypatch, refuse)
    r = env.client.post("/api/v1/ml/analysis", json={})
    assert r.status_code == 503 and r.json()["detail"]["code"] == "langflow_error" and "connect" in r.json()["detail"]["message"]


def test_flow_not_found_is_503(env, monkeypatch):
    _langflow(monkeypatch, lambda r: httpx.Response(404, json={"detail": "Flow not found"}))
    r = env.client.post("/api/v1/ml/analysis", json={})
    assert r.status_code == 503 and "LANGFLOW_FLOW_ID" in r.json()["detail"]["message"]


@pytest.mark.parametrize("code, status", [("target_not_found", 422), ("insufficient_data", 422), ("dataset_unavailable", 503), ("not_enough_models", 502)])
def test_flow_errors_are_mapped(env, monkeypatch, code, status):
    error = {"status": "error", "error": {"code": code, "message": "explained"}}
    _langflow(monkeypatch, lambda r: httpx.Response(200, json=_run_response(json.dumps(error))))
    r = env.client.post("/api/v1/ml/analysis", json={})
    assert r.status_code == status and r.json()["detail"] == {"code": code, "message": "explained"}


@pytest.mark.parametrize("text", ["not json", json.dumps({"status": "success"}), json.dumps([1, 2])])
def test_incomplete_or_malformed_flow_output_is_502(env, monkeypatch, text):
    _langflow(monkeypatch, lambda r: httpx.Response(200, json=_run_response(text)))
    r = env.client.post("/api/v1/ml/analysis", json={})
    assert r.status_code == 502


def test_result_missing_the_ranking_is_rejected(env, monkeypatch):
    broken = {**env.output, "ranking": []}
    _langflow(monkeypatch, lambda r: httpx.Response(200, json=_run_response(json.dumps(broken))))
    r = env.client.post("/api/v1/ml/analysis", json={})
    assert r.status_code == 502 and r.json()["detail"]["code"] == "incomplete_result"


def test_direct_mode_runs_the_same_engine_in_process(env, monkeypatch):
    monkeypatch.setattr(settings, "ML_ANALYSIS_MODE", "direct")
    r = env.client.post("/api/v1/ml/analysis", json={"models": MODELS})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["engine"] == "direct" and body["ranking"][0]["model"] == body["best_model"]["name"]
    best = env.client.get("/api/v1/ml/best-model").json()
    assert best["model_name"] == body["best_model"]["name"] and best["run_id"] == body["run_id"]
    bad = env.client.post("/api/v1/ml/analysis", json={"target": "los_hours"})
    assert bad.status_code == 422 and bad.json()["detail"]["code"] == "target_not_found"


@pytest.mark.parametrize("payload", [{"cv_folds": 1}, {"problem_type": "regression"}, {"unexpected": True}, {"test_size": 0.9}])
def test_request_validation(env, payload):
    assert env.client.post("/api/v1/ml/analysis", json=payload).status_code == 422


def test_requires_simulation_permission(env):
    from app.authentication.dependencies import get_current_user

    env.app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(id=uuid.uuid4(), role=SimpleNamespace(permissions={}), is_active=True)
    assert env.client.post("/api/v1/ml/analysis", json={}).status_code == 403
    assert env.client.get("/api/v1/ml/status").status_code == 200


def test_reports_are_generated_with_interpretations(env, monkeypatch):
    import io

    import openpyxl

    monkeypatch.setattr(settings, "ML_ANALYSIS_MODE", "direct")
    run_id = env.client.post("/api/v1/ml/analysis", json={"models": MODELS}).json()["run_id"]
    pdf = env.client.get(f"/api/v1/ml/analysis/{run_id}/report", params={"format": "pdf"})
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF" and len(pdf.content) > 10_000
    xlsx = env.client.get(f"/api/v1/ml/analysis/{run_id}/report", params={"format": "xlsx"})
    assert xlsx.status_code == 200 and xlsx.content[:2] == b"PK"
    book = openpyxl.load_workbook(io.BytesIO(xlsx.content))
    assert len(book.sheetnames) >= 10
    assert all(str(book[name]["A2"].value).startswith("Interpretación") for name in book.sheetnames[1:])
    assert env.client.get("/api/v1/ml/analysis/20990101T000000Z-deadbeef/report").status_code == 404
    assert env.client.get("/api/v1/ml/analysis/not-a-run-id/report").status_code == 404


def test_status_reports_langflow_reachability(env, monkeypatch):
    _langflow(monkeypatch, lambda r: httpx.Response(200, json={"status": "ok"}))
    assert env.client.get("/api/v1/ml/status").json()["langflow_reachable"] is True
