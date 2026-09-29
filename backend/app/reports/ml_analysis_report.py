"""PDF / XLSX reports of an ML model-comparison analysis (output of the Langflow flow).

Organized by CRISP-DM phase. Every table and chart is followed by its interpretation
(`output["interpretations"]`, generated from the real values by ai_module.ml), so the
report carries the same explanations as the screen. Charts use ReportLab's own graphics
(no matplotlib dependency)."""
from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.graphics.charts.barcharts import HorizontalBarChart
from reportlab.graphics.charts.legends import Legend
from reportlab.graphics.charts.lineplots import LinePlot
from reportlab.graphics.shapes import Drawing, String
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

PAGE = landscape(A4)
MARGIN = 1.2 * cm
WIDTH = PAGE[0] - 2 * MARGIN
PALETTE = [colors.HexColor(c) for c in ("#2563eb", "#16a34a", "#dc2626", "#9333ea", "#ea580c", "#0891b2", "#ca8a04", "#4b5563")]

L = {
    "es": {
        "title": "Análisis comparativo de modelos de Machine Learning", "run": "Ejecución", "engine": "Motor",
        "dataset": "Dataset", "target": "Objetivo", "primary": "Métrica principal", "duration": "Duración (s)",
        "p1": "1. Comprensión del negocio", "p2": "2. Comprensión de los datos (EDA)", "p3": "3. Preparación de los datos",
        "p4": "4. Modelado e hiperparámetros", "p5": "5. Evaluación", "p6": "6. Despliegue y selección",
        "overview": "Resumen del dataset", "missing": "Valores faltantes", "numeric": "Variables numéricas / binarias por clase",
        "categorical": "Variables categóricas", "corr": "Correlaciones de Spearman (pares más fuertes)",
        "prep": "Partición y preprocesamiento", "hyper": "Hiperparámetros seleccionados (búsqueda en rejilla)",
        "cv": "Validación cruzada anidada", "comparison": "Comparación de modelos (conjunto de prueba + CV)",
        "ranking": "Ranking por final_score", "ranking_chart": "Gráfico del ranking (verde: ganador; morado: híbridos)", "roc": "Curvas ROC (conjunto de prueba)", "cm": "Matriz de confusión del mejor modelo",
        "xref": "Referencias cruzadas", "stats": "Pruebas estadísticas inferenciales", "pairwise": "Mejor modelo vs. resto (Holm)",
        "ci": "IC 95% bootstrap (prueba)", "best": "Mejor modelo", "why": "¿Por qué fue seleccionado?", "llm": "Análisis del LLM (no altera la selección)",
        "artifact": "Artefacto persistido", "interp": "Interpretación", "model": "Modelo", "type": "Tipo", "normal": "normal", "hybrid": "híbrido",
        "feature": "Variable", "value": "Valor", "pred": "Predicho", "real": "Real",
    },
    "en": {
        "title": "Machine Learning model comparison", "run": "Run", "engine": "Engine",
        "dataset": "Dataset", "target": "Target", "primary": "Primary metric", "duration": "Duration (s)",
        "p1": "1. Business understanding", "p2": "2. Data understanding (EDA)", "p3": "3. Data preparation",
        "p4": "4. Modeling and hyper-parameters", "p5": "5. Evaluation", "p6": "6. Deployment and selection",
        "overview": "Dataset overview", "missing": "Missing values", "numeric": "Numeric / binary features by class",
        "categorical": "Categorical features", "corr": "Spearman correlations (strongest pairs)",
        "prep": "Split and preprocessing", "hyper": "Selected hyper-parameters (grid search)",
        "cv": "Nested cross-validation", "comparison": "Model comparison (test set + CV)",
        "ranking": "Ranking by final_score", "ranking_chart": "Ranking chart (green: winner; purple: hybrids)", "roc": "ROC curves (test set)", "cm": "Best model confusion matrix",
        "xref": "Cross-reference", "stats": "Inferential statistical tests", "pairwise": "Best model vs. the rest (Holm)",
        "ci": "Bootstrap 95% CI (test)", "best": "Best model", "why": "Why was it selected?", "llm": "LLM analysis (does not change the selection)",
        "artifact": "Persisted artifact", "interp": "Interpretation", "model": "Model", "type": "Type", "normal": "normal", "hybrid": "hybrid",
        "feature": "Feature", "value": "Value", "pred": "Predicted", "real": "Actual",
    },
}


def _f(v: Any, d: int = 3) -> str:
    if v is None:
        return "—"
    if isinstance(v, float):
        return f"{v:.{d}f}"
    return str(v)


# --------------------------------------------------------------------------------- tables (shared by PDF and XLSX)


def report_tables(out: dict[str, Any]) -> list[tuple[str, str, str, pd.DataFrame]]:
    """(phase_key, title_key, interpretation_key, table) in report order."""
    t = L.get(out.get("language", "es"), L["es"])
    eda, it = out["eda"], out["interpretations"]
    ov = eda["overview"]
    tables = [
        ("p2", "overview", "eda_overview", pd.DataFrame([
            (t["dataset"], out["dataset"]["dataset"]), (t["target"], ov["target"]), ("rows", ov["rows"]), ("features", ov["features"]),
            ("positive_rate", ov["positive_rate"]), ("duplicated_rows", ov["duplicated_rows"]), ("complete_rows_pct", ov["complete_rows_pct"]),
            ("groups (patients)", ov["groups"]), ("source", out["dataset"]["source"])], columns=[t["feature"], t["value"]])),
        ("p2", "missing", "eda_missing", pd.DataFrame([m for m in eda["missing"] if m["missing"]] or eda["missing"][:3])),
        ("p2", "numeric", "eda_numeric", pd.DataFrame([{
            t["feature"]: n["feature"], "mean": n["mean"], "median": n["median"], "std": n["std"], "skew": n["skewness"],
            "outliers": n["outliers_iqr"], "mean(+)": n["mean_positive"], "mean(-)": n["mean_negative"],
            "U p (Holm)": n["association"].get("p_value_holm"), "rank-biserial": n["association"]["effect_size_rank_biserial"],
            "sig.": n["association"].get("significant")} for n in eda["numeric"]])),
        ("p2", "categorical", "eda_categorical", pd.DataFrame([{
            t["feature"]: c["feature"], "levels": c["n_levels"], "χ² p (Holm)": c["association"].get("p_value_holm"),
            "Cramér V": c["association"]["cramers_v"], "sig.": c["association"].get("significant"),
            "top level (rate+)": f"{c['levels'][0]['level']} ({_f(c['levels'][0]['positive_rate'])})" if c["levels"] else "—"}
            for c in eda["categorical"]])),
        ("p2", "corr", "eda_correlations", pd.DataFrame(eda["correlations"]["top_pairs"])),
        ("p3", "prep", "preprocessing", pd.DataFrame([(k, v) for k, v in out["preprocessing"]["split"].items()], columns=[t["feature"], t["value"]])),
        ("p4", "hyper", "hyperparameters", pd.DataFrame([{
            t["model"]: name, "grid": h.get("grid_size"), "best params": ", ".join(f"{k}={v}" for k, v in (h.get("best_params") or {}).items()),
            "inner score": h.get("best_inner_score"), "fold stability": h.get("stability")} for name, h in out["hyperparameters"].items()])),
        ("p5", "cv", "cross_validation", pd.DataFrame([{
            t["model"]: m["name"], t["type"]: t[m["family"]], "cv_mean": m["cv_mean"], "cv_std": m["cv_std"],
            "folds": ", ".join(_f(s) for s in m["cv_scores"]), "F1": m["mean"].get("f1"), "PR-AUC": m["mean"].get("pr_auc")}
            for m in out["cross_validation"]["models"].values() if m["status"] == "ok"])),
        ("p5", "comparison", "comparison_table", pd.DataFrame([{
            t["model"]: r["model"], t["type"]: t[r["type"]], "Accuracy": r["accuracy"], "Precision": r["precision"], "Recall": r["recall"],
            "F1": r["f1"], "ROC-AUC": r["roc_auc"], "CV Mean": r["cv_mean"], "CV Std": r["cv_std"], "Overfit gap": r["overfitting_gap"],
            "Final score": next((e["final_score"] for e in out["ranking"] if e["key"] == r["key"]), None)} for r in out["comparison_table"]])),
        ("p5", "ranking", "ranking", pd.DataFrame([{"#": e["rank"], t["model"]: e["model"], t["type"]: t[e["type"]], "final_score": e["final_score"]}
                                                   for e in out["ranking"]])),
        ("p5", "xref", "cross_reference", pd.DataFrame([{
            t["model"]: m["model"], "train": m["train_score"], "cv": m["cv_mean"], "test": m["test_score"], "cv→test": m["cv_vs_test_difference"],
            "stability": m["stability"], "generalization": m["generalization"], "overfitting": m["overfitting"]}
            for m in out["cross_reference_analysis"]["models"].values()])),
        ("p5", "pairwise", "statistical_tests", pd.DataFrame([{
            t["model"]: p["model"], "NB t p(Holm)": p["nadeau_bengio"].get("p_value_holm"), "DeLong p(Holm)": p["delong"].get("p_value_holm"),
            "McNemar p(Holm)": p["mcnemar"].get("p_value_holm"),
            "sig.": any(p[k].get("significant") for k in ("nadeau_bengio", "delong", "mcnemar"))} for p in out["statistical_tests"]["pairwise_vs_best"]])),
        ("p5", "ci", "bootstrap_ci", pd.DataFrame([{t["model"]: name, "ROC-AUC 95% CI": f"[{_f(c['roc_auc'][0])}, {_f(c['roc_auc'][1])}]",
                                                  "F1 95% CI": f"[{_f(c['f1'][0])}, {_f(c['f1'][1])}]"}
                                                 for name, c in out["statistical_tests"]["bootstrap_ci_95"].items()])),
    ]
    return [(p, title, key, df) for p, title, key, df in tables if not df.empty and key in it or key == "eda_overview"]


# --------------------------------------------------------------------------------- PDF


def _table(df: pd.DataFrame, style: ParagraphStyle, header: ParagraphStyle) -> Table:
    data = [[Paragraph(str(c), header) for c in df.columns]]
    data += [[Paragraph(_f(v), style) for v in row] for row in df.itertuples(index=False)]
    table = Table(data, colWidths=[WIDTH / len(df.columns)] * len(df.columns), repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f4f6")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d5db")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    return table


def _ranking_chart(out: dict[str, Any]) -> Drawing:
    ranking = list(reversed(out["ranking"]))
    d = Drawing(WIDTH, 30 + 22 * len(ranking))
    chart = HorizontalBarChart()
    chart.x, chart.y, chart.width, chart.height = 150, 15, WIDTH - 200, 22 * len(ranking)
    chart.data = [[e["final_score"] for e in ranking]]
    chart.categoryAxis.categoryNames = [e["model"] for e in ranking]
    chart.valueAxis.valueMin = 0
    chart.bars[0].fillColor = PALETTE[0]
    for i, e in enumerate(ranking):
        chart.bars[(0, i)].fillColor = PALETTE[1] if e["rank"] == 1 else PALETTE[3] if e["type"] == "hybrid" else PALETTE[0]
    chart.barLabelFormat = "%.4f"
    chart.barLabels.nudge = 18
    d.add(chart)
    return d


def _roc_chart(out: dict[str, Any]) -> Drawing:
    d = Drawing(WIDTH, 260)
    plot = LinePlot()
    plot.x, plot.y, plot.width, plot.height = 60, 30, 330, 210
    names = [n for n, c in out["curves"].items() if c.get("roc")]
    plot.data = [list(zip(out["curves"][n]["roc"]["x"], out["curves"][n]["roc"]["y"])) for n in names] + [[(0, 0), (1, 1)]]
    for i in range(len(names)):
        plot.lines[i].strokeColor = PALETTE[i % len(PALETTE)]
        plot.lines[i].strokeWidth = 2.2 if names[i] == out["best_model"]["name"] else 1
    plot.lines[len(names)].strokeColor = colors.grey
    plot.lines[len(names)].strokeDashArray = [3, 3]
    plot.xValueAxis.valueMin = plot.yValueAxis.valueMin = 0
    plot.xValueAxis.valueMax = plot.yValueAxis.valueMax = 1
    d.add(plot)
    d.add(String(170, 5, "FPR", fontSize=8))
    d.add(String(10, 135, "TPR", fontSize=8))
    legend = Legend()
    legend.x, legend.y = 420, 235
    legend.alignment = "right"  # swatch first, then the model name
    legend.columnMaximum = len(names)  # one column
    legend.colorNamePairs = [(PALETTE[i % len(PALETTE)], n + (" *" if n == out["best_model"]["name"] else "")) for i, n in enumerate(names)]
    legend.fontSize = 8
    legend.dy = legend.dx = 8
    d.add(legend)
    return d


def generate_ml_pdf(out: dict[str, Any], path: Path, engine: str) -> Path:
    t = L.get(out.get("language", "es"), L["es"])
    path.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    cell = ParagraphStyle("c", parent=styles["Normal"], fontSize=7.5, leading=9)
    head = ParagraphStyle("h", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")
    interp = ParagraphStyle("i", parent=styles["Normal"], fontSize=8.5, leading=11, textColor=colors.HexColor("#374151"),
                            backColor=colors.HexColor("#eef2ff"), borderPadding=4, spaceBefore=4, spaceAfter=10)
    h1, h2 = styles["Heading1"], styles["Heading3"]
    it = out["interpretations"]

    def explained(title: str, flowable: Any, key: str | None) -> list[Any]:
        block = [KeepTogether([Paragraph(title, h2), flowable])]
        if key and key in it:
            block.append(Paragraph(f"<b>{t['interp']}:</b> {it[key]}", interp))
        return block

    story: list[Any] = [
        Paragraph(t["title"], styles["Title"]),
        Paragraph(f"{t['run']}: {out['run_id']} · {t['engine']}: {engine} · {t['dataset']}: {out['dataset']['dataset']} · "
                  f"{t['target']}: {out['dataset']['target']} · {t['primary']}: {out['dataset']['primary_metric']} · "
                  f"{t['duration']}: {out.get('duration_s')}", styles["Normal"]),
        Spacer(1, 8),
        Paragraph(t["p1"], h1),
    ]
    story += [Paragraph(f"<b>{p['phase']}</b>: {p['summary']}", styles["Normal"]) for p in out["crisp_dm"]]
    phase = "p1"
    for ph, title, key, df in report_tables(out):
        if ph != phase:
            story += [PageBreak(), Paragraph(t[ph], h1)]
            phase = ph
        if title == "ranking":  # table and chart share one interpretation, shown once after both
            story += explained(t[title], _table(df, cell, head), None)
            story += explained(t["ranking_chart"], _ranking_chart(out), key)
            continue
        story += explained(t[title], _table(df, cell, head), key)
        if title == "ci":
            story += explained(t["roc"], _roc_chart(out), "roc_curves")
    best = out["best_model"]
    (tn, fp), (fn, tp) = best["metrics"]["confusion_matrix"]
    cm_df = pd.DataFrame([[f"{t['real']} 0", tn, fp], [f"{t['real']} 1", fn, tp]], columns=["", f"{t['pred']} 0", f"{t['pred']} 1"])
    story += [PageBreak(), Paragraph(t["p6"], h1)]
    story += explained(t["cm"], _table(cm_df, cell, head), "confusion_matrix")
    story += [Paragraph(f"{t['best']}: <b>{best['name']}</b> ({t[best['type']]}) · final_score {best['final_score']:.4f}", styles["Heading2"]),
              Paragraph(f"<b>{t['why']}</b> {out['selection_reason']}", interp)]
    if out.get("analysis_source") == "llm":
        story.append(Paragraph(f"<b>{t['llm']}:</b> {out['analysis']}", interp))
    story.append(Paragraph(f"{t['artifact']}: {best['artifact_path']}", cell))
    SimpleDocTemplate(str(path), pagesize=PAGE, leftMargin=MARGIN, rightMargin=MARGIN, topMargin=MARGIN, bottomMargin=MARGIN,
                      title=t["title"]).build(story)
    return path


# --------------------------------------------------------------------------------- XLSX

_INVALID_SHEET_CHARS = set("[]:*?/\\")  # Excel rejects these in sheet names


def generate_ml_xlsx(out: dict[str, Any], path: Path) -> Path:
    t = L.get(out.get("language", "es"), L["es"])
    path.parent.mkdir(parents=True, exist_ok=True)
    it = out["interpretations"]
    used: set[str] = set()
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary = pd.DataFrame([(p["phase"], p["summary"]) for p in out["crisp_dm"]] +
                               [(t["best"], out["best_model"]["name"]), ("final_score", out["best_model"]["final_score"]),
                                (t["why"], out["selection_reason"])], columns=["CRISP-DM", ""])
        summary.to_excel(writer, sheet_name="Resumen" if t is L["es"] else "Summary", index=False)
        for _, title, key, df in report_tables(out):
            name = "".join(ch for ch in t[title] if ch not in _INVALID_SHEET_CHARS).strip()[:31]
            while name in used:
                name = name[:29] + "_" + str(len(used))
            used.add(name)
            df.to_excel(writer, sheet_name=name, index=False, startrow=3)
            sheet = writer.sheets[name]
            sheet["A1"] = t[title]
            sheet["A1"].font = Font(bold=True, size=12)
            sheet["A2"] = f"{t['interp']}: {it[key]}"
            sheet["A2"].alignment = Alignment(wrap_text=True, vertical="top")
            sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(len(df.columns), 6))
            sheet.row_dimensions[2].height = 90
            for cell in sheet[4]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F2937")
            for column in sheet.columns:
                sheet.column_dimensions[column[3].column_letter].width = 18
    return path
