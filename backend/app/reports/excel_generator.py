"""XLSX export of historical data, simulation results, and KPIs (pandas + openpyxl)."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill

from app.reports.labels import STATUS_COLORS_HEX, STATUS_LABELS_ES, column_label


def generate_excel_report(
    output_path: Path,
    summary_metrics: dict,
    timeseries: list[dict],
    patients: list[dict] | None = None,
) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        pd.DataFrame(list(summary_metrics.items()), columns=["Métrica", "Valor"]).to_excel(
            writer, sheet_name="Indicadores", index=False
        )

        # Column headers are translated, but values keep their native int/float
        # type (unlike the DOCX/PDF text tables) so the sheet stays usable for
        # sorting/filtering/formulas — only the categorical status is mapped.
        ts_df = pd.DataFrame(
            [
                {
                    column_label(c): STATUS_LABELS_ES.get(v, v) if c == "status_band" else v
                    for c, v in row.items()
                }
                for row in timeseries
            ]
        )
        ts_df.to_excel(writer, sheet_name="ResultadosSimulacion", index=False)

        if timeseries and "status_band" in timeseries[0]:
            ws = writer.sheets["ResultadosSimulacion"]
            status_col_idx = list(ts_df.columns).index(column_label("status_band")) + 1  # openpyxl is 1-based
            for row_i, row in enumerate(timeseries, start=2):  # row 1 is the header
                hex_color = STATUS_COLORS_HEX.get(row.get("status_band"))
                if not hex_color:
                    continue
                cell = ws.cell(row=row_i, column=status_col_idx)
                cell.fill = PatternFill(start_color=f"FF{hex_color.upper()}", end_color=f"FF{hex_color.upper()}", fill_type="solid")
                cell.font = Font(color="FFFFFFFF", bold=True)

        if patients:
            pd.DataFrame(patients).to_excel(writer, sheet_name="Pacientes", index=False)

    return output_path
