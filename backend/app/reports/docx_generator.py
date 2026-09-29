"""DOCX executive/scientific report generation (python-docx)."""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.reports.labels import STATUS_COLORS_HEX, column_label, format_value


def _shade_cell(cell, hex_color: str) -> None:
    """python-docx has no cell-background API — this sets the underlying
    OOXML shading element directly."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def generate_docx_report(
    output_path: Path,
    hospital_name: str,
    event_name: str,
    simulation_id: str,
    summary_metrics: dict,
    timeseries_sample: list[dict],
    ai_recommendations: list[str],
) -> Path:
    doc = Document()

    # Cover page
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Hospital Digital Twin\nReporte de Simulación — Evento de Víctimas Masivas")
    run.bold = True
    run.font.size = Pt(24)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(f"{hospital_name}\nEvento: {event_name}\nID de Simulación: {simulation_id}\n"
                f"Generado: {datetime.now(timezone.utc).isoformat()} UTC")
    doc.add_page_break()

    # Executive summary
    doc.add_heading("Resumen Ejecutivo", level=1)
    doc.add_paragraph(
        "Este reporte documenta una simulación híbrida de Agentes / Dinámica de Sistemas sobre "
        "la asignación de recursos hospitalarios ante un escenario de víctimas masivas. Los "
        "resultados a continuación resumen el flujo de pacientes, la saturación de recursos y "
        "las recomendaciones de reasignación generadas por IA."
    )

    doc.add_heading("Resultados Clave", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text = "Métrica", "Valor"
    for key, value in summary_metrics.items():
        row = table.add_row().cells
        row[0].text = str(key)
        row[1].text = str(value)

    doc.add_heading("Serie Temporal de la Simulación (muestra)", level=1)
    if timeseries_sample:
        cols = list(timeseries_sample[0].keys())
        ts_table = doc.add_table(rows=1, cols=len(cols))
        ts_table.style = "Light List Accent 1"
        for i, c in enumerate(cols):
            ts_table.rows[0].cells[i].text = column_label(c)
        for row_data in timeseries_sample[:30]:
            row = ts_table.add_row().cells
            for i, c in enumerate(cols):
                raw = row_data.get(c, "")
                cell = row[i]
                cell.text = format_value(c, raw)
                if c == "status_band" and raw in STATUS_COLORS_HEX:
                    _shade_cell(cell, STATUS_COLORS_HEX[raw])
                    run = cell.paragraphs[0].runs[0]
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    doc.add_heading("Recomendaciones Generadas por IA", level=1)
    for rec in ai_recommendations or ["No se generaron recomendaciones para esta ejecución."]:
        doc.add_paragraph(rec, style="List Bullet")

    doc.add_heading("Nota Metodológica", level=1)
    doc.add_paragraph(
        "Los modelos predictivos y de optimización fueron entrenados/validados combinando "
        "conjuntos de datos públicos de salud procesados (ETL) y trazas sintéticas de eventos de "
        "víctimas masivas generadas por este mismo motor de simulación híbrido, según la nota de "
        "reproducibilidad en docs/ARCHITECTURE.md sección 7."
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path
